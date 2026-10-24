"""PaperForge desktop conversion engine. All document processing stays on this computer."""
from __future__ import annotations
import argparse, csv, hashlib, importlib.metadata, io, json, os, shutil, subprocess, sys
import tempfile, time, uuid, zipfile
from pathlib import Path

APP_ROOT = Path(sys.executable).parent if getattr(sys,"frozen",False) else Path(__file__).parent

OFFICE_FAMILIES = {
    'writer': {'doc','docx','odt','rtf','txt','html','htm','md'},
    'calc': {'xls','xlsx','ods','csv','tsv'},
    'impress': {'ppt','pptx','odp'},
    'draw': {'svg','odg'},
}
OFFICE_TARGETS = {'pdf': None, 'docx':'writer', 'odt':'writer', 'rtf':'writer',
    'xlsx':'calc', 'ods':'calc', 'csv':'calc', 'pptx':'impress', 'odp':'impress'}

def sha256(path):
    with Path(path).open('rb') as stream: return hashlib.file_digest(stream, 'sha256').hexdigest()

def config():
    path = Path(os.environ.get('PAPERFORGE_DESKTOP_CONFIG', APP_ROOT/'config.json'))
    return json.loads(path.read_text(encoding='utf-8-sig')) if path.exists() else {}

def soffice():
    settings = config()
    candidates = [os.environ.get('PAPERFORGE_SOFFICE'), settings.get('soffice'), shutil.which('soffice'),
        APP_ROOT/'runtime/LibreOffice/program/soffice.com',
        Path(os.environ.get('PROGRAMFILES', 'C:/Program Files'))/'LibreOffice/program/soffice.com',
        Path('/Applications/LibreOffice.app/Contents/MacOS/soffice')]
    return next((str(p) for p in candidates if p and Path(p).is_file()), None)

def capabilities():
    versions = {}
    for package in ['pymupdf','pdf2docx','openpyxl','python-pptx','rapidocr-onnxruntime','opencv-python-headless']:
        try: versions[package] = importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError: versions[package] = None
    executable = soffice()
    lo_version = None
    if executable:
        result = subprocess.run([executable,'--version'],capture_output=True,text=True,timeout=30,
            creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
        lo_version = result.stdout.strip()
    return {'libreoffice':lo_version,'soffice':executable,'packages':versions,'networkRequired':False}

def office(source, target, work, options):
    extension = source.suffix.lower().lstrip('.')
    family = next((name for name, formats in OFFICE_FAMILIES.items() if extension in formats), None)
    if family is None: raise ValueError('此 Office 输入格式尚未接入：'+extension)
    if target not in OFFICE_TARGETS: raise ValueError('目标格式不支持：'+target)
    expected = OFFICE_TARGETS[target]
    if expected and expected != family: raise ValueError('请在同类 Office 格式间转换；跨类型请使用 PDF')
    executable = soffice()
    if not executable: raise RuntimeError('未找到 LibreOffice。请设置 config.json 的 soffice 路径或安装 LibreOffice。')
    profile = work/'profile'
    (profile/'user').mkdir(parents=True)
    # Per-job profile prevents a running desktop Office process from swallowing the conversion.
    (profile/'user/registrymodifications.xcu').write_text('''<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item></oor:items>''',encoding='utf-8')
    export = target
    if target == 'pdf':
        filter_name = {'writer':'writer_pdf_Export','calc':'calc_pdf_Export','impress':'impress_pdf_Export','draw':'draw_pdf_Export'}[family]
        properties = {'ExportBookmarks':{'type':'boolean','value':'true'}, 'ExportNotesPages':{'type':'boolean','value':'false'}}
        if family == 'calc' and options.get('fitWidth', False):
            properties['SinglePageSheets'] = {'type':'boolean','value':'true'}
        export += ':'+filter_name+':'+json.dumps(properties,separators=(',',':'))
    output_dir = work/'converted'; output_dir.mkdir()
    command = [executable,'-env:UserInstallation='+profile.as_uri(),'--headless','--nologo','--norestore',
        '--convert-to',export,'--outdir',str(output_dir),str(source)]
    result = subprocess.run(command,capture_output=True,text=True,encoding='utf-8',errors='replace',
        timeout=int(options.get('timeout',300)),creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
    outputs = list(output_dir.glob('*.'+target))
    if result.returncode or not outputs: raise RuntimeError('LibreOffice 转换失败：'+(result.stderr or result.stdout)[-1800:])
    return outputs, {'engine':'LibreOffice','log':result.stdout.strip().replace(str(work),'<job>'),
        'notes':['使用实际 Office 排版引擎。字体替代、修订、浮动对象和公式仍需检查。',
                 'PDF 保留视觉页面；导出 CSV 只保留一个工作表，不包含样式与公式。']}

_OCR = None
def ocr_engine():
    global _OCR
    if _OCR is None:
        from rapidocr_onnxruntime import RapidOCR
        _OCR = RapidOCR()
    return _OCR

def raster_lines(page):
    import cv2, numpy as np, pymupdf
    pix = page.get_pixmap(matrix=pymupdf.Matrix(2,2),alpha=False)
    image = np.frombuffer(pix.samples,np.uint8).reshape(pix.height,pix.width,3)
    gray = cv2.cvtColor(image,cv2.COLOR_RGB2GRAY)
    binary = cv2.adaptiveThreshold(gray,255,cv2.ADAPTIVE_THRESH_GAUSSIAN_C,cv2.THRESH_BINARY_INV,31,15)
    lines=[]
    for horizontal in (True,False):
        kernel=cv2.getStructuringElement(cv2.MORPH_RECT,(max(20,pix.width//35),1) if horizontal else (1,max(20,pix.height//35)))
        mask=cv2.morphologyEx(binary,cv2.MORPH_OPEN,kernel)
        contours,_=cv2.findContours(mask,cv2.RETR_LIST,cv2.CHAIN_APPROX_SIMPLE)
        for contour in contours:
            x,y,w,h=cv2.boundingRect(contour)
            if horizontal and w>=max(25,pix.width//30): lines.append(((x/2,(y+h/2)/2),((x+w)/2,(y+h/2)/2)))
            if not horizontal and h>=max(25,pix.height//30): lines.append((((x+w/2)/2,y/2),((x+w/2)/2,(y+h)/2)))
    return lines,image

def add_ocr(page, image, blocks=None):
    import pymupdf
    results,_=ocr_engine()(image)
    count=0
    for box,text,score in results or []:
        if not text.strip():continue
        x=min(p[0] for p in box)/2; y=min(p[1] for p in box)/2
        w=(max(p[0] for p in box)-min(p[0] for p in box))/2
        h=(max(p[1] for p in box)-min(p[1] for p in box))/2
        if blocks is not None: blocks.append({"text":text,"bbox":(x,y,x+w,y+h)})
        font=pymupdf.Font('china-s')
        natural=max(1,font.text_length(text,fontsize=1))
        size=min(max(3,h*.8),max(3,w/natural))
        page.insert_text((x,y+h*.86),text,fontsize=size,fontname='china-s',render_mode=3)
        count+=1
    return count

def open_pdf(source):
    import pymupdf
    if source.suffix.lower()=='.pdf':return pymupdf.open(source)
    from PIL import Image,ImageOps,ImageSequence
    result=pymupdf.open()
    with Image.open(source) as image:
        for frame in ImageSequence.Iterator(image):
            frame=ImageOps.exif_transpose(frame).convert('RGB')
            data=io.BytesIO();frame.save(data,format='PNG')
            page=result.new_page(width=frame.width*.5,height=frame.height*.5)
            page.insert_image(page.rect,stream=data.getvalue())
    return result

def table_records(source, options, progress):
    records=[];ocr_pages=[]
    with open_pdf(source) as pdf:
        for number,page in enumerate(pdf):
            progress(f'分析表格 {number+1} / {len(pdf)}')
            page.set_rotation(0)
            scanned=len(page.get_text().strip())<12
            lines=[];ocr_blocks=[]
            if scanned and options.get('ocr',True):
                lines,image=raster_lines(page)
                add_ocr(page,image,ocr_blocks);ocr_pages.append(number+1)
            strategy=options.get('tableStrategy','auto')
            methods=['lines_strict','text'] if strategy=='auto' else [strategy]
            found=[]
            for method in methods:
                tables=page.find_tables(strategy=method,add_lines=lines or None).tables
                tables=[t for t in tables if t.row_count>=2 and t.col_count>=2]
                if tables:found=tables;break
            for table in found:
                def cell_text(cell):
                    if cell is None:return None
                    if not ocr_blocks:return page.get_textbox(cell).strip()
                    x0,y0,x1,y1=cell
                    selected=[b for b in ocr_blocks if x0<=(b['bbox'][0]+b['bbox'][2])/2<x1 and y0<=(b['bbox'][1]+b['bbox'][3])/2<y1]
                    selected.sort(key=lambda b:(round(b['bbox'][1]/4),b['bbox'][0]))
                    return '\n'.join(b['text'] for b in selected)
                data=[[cell_text(cell) for cell in row.cells] for row in table.rows]
                # Own all cell data while the page/table finder is alive.
                records.append({'page':number+1,'bbox':list(table.bbox),'rows':data,
                    'cells':[list(c) if c else None for row in table.rows for c in row.cells],
                    'rowCount':table.row_count,'colCount':table.col_count,'strategy':method,'ocr':scanned})
    if not records:raise ValueError('未找到可重建表格。请裁切到表格区域，或尝试有清晰边框/对齐列的原件。')
    return records,ocr_pages

def cell_merges(table):
    boxes=[c for c in table['cells'] if c]
    xs=sorted({round(c[i],1) for c in boxes for i in [0,2]})
    ys=sorted({round(c[i],1) for c in boxes for i in [1,3]})
    spans=[]
    for box in boxes:
        x0,y0,x1,y1=[round(v,1) for v in box]
        c0,c1=xs.index(x0),xs.index(x1)-1
        r0,r1=ys.index(y0),ys.index(y1)-1
        if (c1>c0 or r1>r0) and r1<table['rowCount'] and c1<table['colCount']:
            spans.append([r0,c0,r1,c1])
    return spans

def join_tables(records, enabled):
    joined=[]
    for record in records:
        record['merges']=cell_merges(record)
        record['sourcePages']=[record['page']]
        previous=joined[-1] if joined else None
        if enabled and previous and previous['sourcePages'][-1]+1==record['page'] and previous['colCount']==record['colCount'] and previous['rows'][0]==record['rows'][0]:
            offset=len(previous['rows'])-1
            previous['rows'].extend(record['rows'][1:])
            previous['merges'].extend([[r0+offset,c0,r1+offset,c1] for r0,c0,r1,c1 in record['merges'] if r0>=1])
            previous['rowCount']=len(previous['rows'])
            previous['sourcePages'].append(record['page'])
        else:joined.append(record)
    return joined

def tables(source, kind, work, options, progress):
    records,ocr_pages=table_records(source,options,progress)
    records=join_tables(records,options.get('joinTables',False))
    output=work/('tables.'+('xlsx' if kind=='xlsx' else 'zip'))
    if kind=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font,Alignment,PatternFill,Border,Side
        from openpyxl.utils import get_column_letter
        book=Workbook();book.remove(book.active)
        for n,table in enumerate(records):
            sheet=book.create_sheet(f'P{table["page"]}-T{n+1}')
            for row in table['rows']:
                sheet.append(['' if value is None else str(value) for value in row])
            # Keep identifiers and strings literally, including leading zeroes and formula-like text.
            for row in sheet:
                for cell in row:
                    cell.data_type='s';cell.number_format='@';cell.alignment=Alignment(vertical='top',wrap_text=True)
                    cell.border=Border(*( [Side(style='thin',color='D6DDD8')]*4 ))
            for cell in sheet[1]:cell.font=Font(bold=True);cell.fill=PatternFill('solid',fgColor='E1EFE6')
            for r0,c0,r1,c1 in table['merges']:
                sheet.merge_cells(start_row=r0+1,start_column=c0+1,end_row=r1+1,end_column=c1+1)
            for c in range(1,table['colCount']+1):sheet.column_dimensions[get_column_letter(c)].width=22
            sheet.freeze_panes='A2';sheet.sheet_view.showGridLines=False
        info=book.create_sheet('来源与说明')
        info.append(['表格','源页码','结构策略','是否OCR'])
        for n,t in enumerate(records):info.append([n+1,','.join(map(str,t['sourcePages'])),t['strategy'],str(t['ocr'])])
        book.save(output)
    else:
        with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as archive:
            for n,t in enumerate(records):
                text=io.StringIO(newline='');writer=csv.writer(text)
                writer.writerows(t['rows'])
                archive.writestr(f'page-{t["page"]}-table-{n+1}.csv','\ufeff'+text.getvalue())
            archive.writestr('structure.json',json.dumps(records,ensure_ascii=False,indent=2))
    return [output],{'engine':'PyMuPDF + OpenCV + RapidOCR','tableCount':len(records),'ocrPages':ocr_pages,
        'notes':['XLSX 保留文本、行列与识别到的合并单元格；CSV 不表达合并关系。',
                 '无线表按对齐位置推断。开启跨页合并时，仅连接相邻页、同列数、同首行表头的表格，并保留来源页码。',
                 'OCR 与缺线表存在识别误差，原件和结构清单可用于对照。']}

def pdf_docx(source,work,options,progress):
    from pdf2docx import Converter
    import pymupdf
    input_pdf=source;ocr_count=0
    if options.get('ocr',True):
        with pymupdf.open(source) as document:
            for page in document:
                if len(page.get_text().strip())<12:
                    _,image=raster_lines(page);ocr_count+=add_ocr(page,image)
            if ocr_count:
                input_pdf=work/'ocr-input.pdf';document.save(input_pdf)
    output=work/'document.docx';progress('重建 Word 版面、图片与表格')
    converter=Converter(str(input_pdf))
    try:converter.convert(str(output),multi_processing=False)
    finally:converter.close()
    return [output],{'engine':'pdf2docx','ocrBlocks':ocr_count,
        'notes':['输出可编辑段落、图片和识别到的表格。复杂脚注、跨栏与字体可能重新排版。',
                 'PDF 缺少原始 Office 结构，重建结果需对照原件；扫描页的原图可能作为背景保留。']}

def pdf_pptx(source,work,options,progress):
    import pymupdf
    from pptx import Presentation
    from pptx.util import Pt
    output=work/'presentation.pptx';presentation=Presentation()
    with pymupdf.open(source) as document:
        first=document[0].rect
        presentation.slide_width=Pt(first.width);presentation.slide_height=Pt(first.height)
        for n,page in enumerate(document):
            progress(f'生成幻灯片 {n+1} / {len(document)}')
            slide=presentation.slides.add_slide(presentation.slide_layouts[6])
            image=io.BytesIO(page.get_pixmap(matrix=pymupdf.Matrix(2,2),alpha=False).tobytes('png'))
            ratio=min(first.width/page.rect.width,first.height/page.rect.height)
            w,h=page.rect.width*ratio,page.rect.height*ratio
            slide.shapes.add_picture(image,Pt((first.width-w)/2),Pt((first.height-h)/2),width=Pt(w),height=Pt(h))
    presentation.save(output)
    return [output],{'engine':'PyMuPDF + python-pptx','notes':['整页图像幻灯片，保留页面外观；文字与图表不是独立可编辑对象。']}

def convert(source,mode,work,options=None,progress=print):
    options=options or {}
    if mode.startswith('office-'):return office(source,mode[7:],work,options)
    if mode=='pdf-docx':return pdf_docx(source,work,options,progress)
    if mode in ('tables-xlsx','tables-csv'):return tables(source,mode.split('-')[1],work,options,progress)
    if mode=='pdf-pptx':return pdf_pptx(source,work,options,progress)
    raise ValueError('不支持的转换任务：'+mode)

def create_task(source, mode, destination, options=None):
    source=Path(source).resolve();destination=Path(destination).resolve();destination.mkdir(parents=True,exist_ok=True)
    output=destination/('PaperForge-task-'+uuid.uuid4().hex[:12]+'.zip')
    name='input'+source.suffix.lower()
    with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as archive:
        archive.write(source,name)
        archive.writestr('task.json',json.dumps({'format':'PaperForge-task-1','id':uuid.uuid4().hex,
            'mode':mode,'input':name,'originalName':source.name,'sha256':sha256(source),
            'options':options or {'ocr':True,'tableStrategy':'auto'}},ensure_ascii=False,indent=2))
    return output

def process_task(package, destination, progress=print):
    package=Path(package).resolve();destination=Path(destination).resolve();destination.mkdir(parents=True,exist_ok=True)
    task_id=uuid.uuid4().hex
    with tempfile.TemporaryDirectory(prefix='paperforge-',dir=destination) as temporary:
        work=Path(temporary)
        with zipfile.ZipFile(package) as archive:
            manifest=archive.getinfo('task.json')
            if manifest.file_size>1024*1024:raise ValueError('任务清单过大')
            task=json.loads(archive.read(manifest))
            if task.get('format')!='PaperForge-task-1':raise ValueError('任务包版本不受支持')
            name=task['input']
            if name!=Path(name).name or '/' in name or '\\' in name:raise ValueError('输入文件路径无效')
            entry=archive.getinfo(name)
            if not 0<entry.file_size<=1024**3:raise ValueError('输入文件为空或超过1GB')
            source=work/name
            with archive.open(entry) as stream,source.open('wb') as target:shutil.copyfileobj(stream,target)
            if sha256(source)!=task['sha256']:raise ValueError('输入文件 SHA256 校验失败')
        progress('任务已读取：'+task.get('originalName',name))
        started=time.time();outputs,report=convert(source,task['mode'],work,task.get('options',{}),progress)
        result=destination/('PaperForge-result-'+task_id[:12]+'.zip')
        records=[]
        with zipfile.ZipFile(result,'w',zipfile.ZIP_DEFLATED) as archive:
            for output in outputs:
                if not output.is_file() or output.stat().st_size==0:raise RuntimeError('转换结果为空')
                records.append({'path':output.name,'bytes':output.stat().st_size,'sha256':sha256(output)})
                archive.write(output,output.name)
            result_manifest={'format':'PaperForge-result-1','taskId':task.get('id'),
                'originalName':task.get('originalName'),'mode':task['mode'],'outputs':records,
                'report':report,'elapsedSeconds':round(time.time()-started,2),'engines':{key:value for key,value in capabilities().items() if key!='soffice'}}
            archive.writestr('manifest.json',json.dumps(result_manifest,ensure_ascii=False,indent=2))
        progress('完成：'+str(result))
        return result,result_manifest

def main():
    parser=argparse.ArgumentParser(description='PaperForge 离线电脑助手')
    parser.add_argument('task',nargs='?',type=Path)
    parser.add_argument('--output',type=Path,default=Path.cwd()/'PaperForge-results')
    parser.add_argument('--capabilities',action='store_true')
    parser.add_argument('--mode',default='office-pdf',help='直接处理电脑文件时的转换方式')
    args=parser.parse_args()
    if args.capabilities:print(json.dumps(capabilities(),ensure_ascii=False,indent=2));return
    if args.task:
        package=args.task if args.task.suffix.lower()=='.zip' else create_task(args.task,args.mode,args.output)
        process_task(package,args.output)
    else:
        from gui import run
        run()
if __name__=='__main__':main()
